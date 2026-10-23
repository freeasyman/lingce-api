#!/usr/bin/env python3
"""周报/月报导出：从 JSON 报告数据生成 Word(docx) + PPT(pptx)。

用法：python3 report_export.py < report.json
输出：stdout 打印 {word, ppt} 文件路径。
"""
import json
import sys
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt


def load_report():
    return json.load(sys.stdin)


def gen_word(report, out_path):
    doc = Document()
    doc.add_heading('内容运营周报/月报', level=0)

    total = report.get('total', 0)
    by_account = report.get('by_account', {}) or {}
    by_goal = report.get('by_goal', {}) or {}
    items = report.get('items', []) or []

    # 1. 本月亮点
    doc.add_heading('一、本月亮点', level=1)
    accounts = '、'.join(by_account.keys()) or '暂无'
    doc.add_paragraph(f'本月发布 {total} 篇，覆盖 {len(by_account)} 个账号（{accounts}），覆盖 {len(by_goal)} 个内容方向。')

    # 2. 数据总览（空字段让填）
    doc.add_heading('二、数据总览（请补充后台数据）', level=1)
    table = doc.add_table(rows=2, cols=6)
    headers = ['阅读量', '播放量', '点赞', '评论', '涨粉', '咨询线索']
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        table.rows[1].cells[i].text = '____'

    # 3. 内容矩阵
    doc.add_heading('三、内容矩阵', level=1)
    goal_cn = {'science': '科普', 'brand': '品牌', 'doctor_ip': '医生IP', 'conversion': '转化', 'geo': 'GEO'}
    matrix = '、'.join(f'{goal_cn.get(g, g)} {n} 篇' for g, n in by_goal.items()) or '暂无'
    doc.add_paragraph(matrix)

    # 4. 专业亮点
    doc.add_heading('四、专业亮点', level=1)
    doc.add_paragraph('· 补 AI 收录缺口：____ 个\n· 命中搜索词：____ 个\n· 覆盖患者高频问题：____ 个')

    # 5. 计划完成
    doc.add_heading('五、计划完成', level=1)
    doc.add_paragraph(f'本月完成 {total} 篇。')

    # 6. 内容明细
    doc.add_heading('六、内容明细', level=1)
    t2 = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(['标题', '账号', '形式', '阅读', '点赞']):
        t2.rows[0].cells[i].text = h
    for it in items:
        row = t2.add_row().cells
        row[0].text = it.get('title', '')
        row[1].text = it.get('account', '')
        row[2].text = it.get('form', '')
        row[3].text = str(it.get('read_count', '') or '')
        row[4].text = str(it.get('like_count', '') or '')

    doc.save(out_path)


def gen_ppt(report, out_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    total = report.get('total', 0)
    by_account = report.get('by_account', {}) or {}
    by_goal = report.get('by_goal', {}) or {}

    # 封面
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = box.text_frame
    tf.text = f'内容运营周报/月报'
    p = tf.add_paragraph()
    p.text = f'本月发布 {total} 篇 · 覆盖 {len(by_account)} 账号 · {len(by_goal)} 方向'

    # 数据总览
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(5))
    tf = box.text_frame
    tf.text = '数据总览'
    for k, v in by_account.items():
        p = tf.add_paragraph()
        p.text = f'{k}：{v} 篇'

    # 内容矩阵
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(5))
    tf = box.text_frame
    tf.text = '内容矩阵'
    goal_cn = {'science': '科普', 'brand': '品牌', 'doctor_ip': '医生IP', 'conversion': '转化', 'geo': 'GEO'}
    for g, n in by_goal.items():
        p = tf.add_paragraph()
        p.text = f'{goal_cn.get(g, g)}：{n} 篇'

    prs.save(out_path)


def main():
    report = load_report()
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    word_path = f'/tmp/content_report_{ts}.docx'
    ppt_path = f'/tmp/content_report_{ts}.pptx'
    gen_word(report, word_path)
    gen_ppt(report, ppt_path)
    print(json.dumps({'word': word_path, 'ppt': ppt_path}))


if __name__ == '__main__':
    main()
